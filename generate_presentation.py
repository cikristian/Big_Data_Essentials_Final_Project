from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_CONNECTOR, MSO_SHAPE
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.util import Inches, Pt


OUT = Path(__file__).resolve().parent / "Kigali_Transport_Delay_Analytics.pptx"
prs = Presentation()
prs.slide_width = Inches(13.333)
prs.slide_height = Inches(7.5)

NAVY = RGBColor(10, 22, 40)
NAVY_2 = RGBColor(16, 35, 63)
CYAN = RGBColor(66, 197, 230)
BLUE = RGBColor(8, 124, 240)
ORANGE = RGBColor(242, 140, 54)
MINT = RGBColor(30, 185, 128)
WHITE = RGBColor(247, 250, 253)
MUTED = RGBColor(163, 180, 198)
LINE = RGBColor(47, 69, 96)


def box(slide, x, y, w, h, fill=NAVY_2, line=None, radius=True):
    shape = slide.shapes.add_shape(
        MSO_SHAPE.ROUNDED_RECTANGLE if radius else MSO_SHAPE.RECTANGLE,
        Inches(x), Inches(y), Inches(w), Inches(h),
    )
    shape.fill.solid()
    shape.fill.fore_color.rgb = fill
    shape.line.color.rgb = line or fill
    shape.line.width = Pt(1)
    return shape


def text(slide, value, x, y, w, h, size=18, color=WHITE, bold=False,
         font="Aptos", align=PP_ALIGN.LEFT, valign=MSO_ANCHOR.TOP):
    tx = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = tx.text_frame
    tf.clear()
    tf.word_wrap = True
    tf.vertical_anchor = valign
    p = tf.paragraphs[0]
    p.alignment = align
    run = p.add_run()
    run.text = value
    run.font.name = font
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = color
    return tx


def bullets(slide, items, x, y, w, h, size=16, color=WHITE, accent=CYAN):
    tx = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = tx.text_frame
    tf.clear()
    tf.word_wrap = True
    for index, item in enumerate(items):
        p = tf.paragraphs[0] if index == 0 else tf.add_paragraph()
        p.text = f"- {item}"
        p.level = 0
        p.space_after = Pt(10)
        p.font.name = "Aptos"
        p.font.size = Pt(size)
        p.font.color.rgb = color
    return tx


def base(title, kicker, number):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    slide.background.fill.solid()
    slide.background.fill.fore_color.rgb = NAVY
    box(slide, 0, 0, 0.16, 7.5, CYAN, CYAN, False)
    text(slide, kicker.upper(), 0.65, 0.38, 3.7, 0.25, 9, CYAN, True)
    text(slide, title, 0.65, 0.72, 11.5, 0.55, 27, WHITE, True, "Aptos Display")
    text(slide, f"{number:02d}", 12.25, 0.42, 0.45, 0.3, 11, MUTED, True, align=PP_ALIGN.RIGHT)
    box(slide, 0.65, 1.42, 12.0, 0.012, LINE, LINE, False)
    text(slide, "KIGALI PUBLIC TRANSPORT DELAY ANALYTICS", 0.65, 7.12, 5.5, 0.18, 8, MUTED, True)
    return slide


def arrow(slide, x1, y1, x2, y2, color=CYAN):
    line = slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, Inches(x1), Inches(y1), Inches(x2), Inches(y2))
    line.line.color.rgb = color
    line.line.width = Pt(2)
    line.line.end_arrowhead = True
    return line


def add_kpi(slide, value, label, x, y, color=CYAN):
    box(slide, x, y, 2.7, 1.15, NAVY_2, LINE)
    text(slide, value, x + 0.18, y + 0.17, 2.3, 0.42, 25, color, True, "Aptos Display")
    text(slide, label.upper(), x + 0.18, y + 0.72, 2.3, 0.2, 9, MUTED, True)


# 1. Cover
slide = prs.slides.add_slide(prs.slide_layouts[6])
slide.background.fill.solid(); slide.background.fill.fore_color.rgb = NAVY
box(slide, 0, 0, 0.2, 7.5, CYAN, CYAN, False)
box(slide, 8.85, 0, 4.48, 7.5, NAVY_2, NAVY_2, False)
text(slide, "BIG DATA ESSENTIALS · FINAL PROJECT", 0.78, 0.72, 5.7, 0.25, 10, CYAN, True)
text(slide, "Kigali Public\nTransport Delay\nAnalytics", 0.78, 1.55, 7.5, 2.15, 34, WHITE, True, "Aptos Display")
text(slide, "A live data pipeline for generating, streaming, storing, and interpreting public transport events in near-real time.", 0.82, 4.2, 6.6, 0.75, 17, MUTED)
box(slide, 0.82, 5.55, 2.1, 0.42, BLUE, BLUE)
text(slide, "RWANDA · KIGALI", 0.98, 5.66, 1.8, 0.16, 9, WHITE, True, align=PP_ALIGN.CENTER)
text(slide, "Django  ·  Kafka  ·  MySQL  ·  HDFS  ·  Spark MLlib", 0.82, 6.72, 7.1, 0.22, 11, MUTED)
for i, color in enumerate([CYAN, BLUE, MINT, ORANGE]):
    box(slide, 9.55, 1.15 + i * 1.15, 2.1, 0.72, color, color)
    text(slide, ["GENERATE", "STREAM", "STORE", "PREDICT"][i], 9.55, 1.38 + i * 1.15, 2.1, 0.2, 12, NAVY, True, align=PP_ALIGN.CENTER)
arrow(slide, 10.6, 1.88, 10.6, 2.3, WHITE); arrow(slide, 10.6, 3.03, 10.6, 3.45, WHITE); arrow(slide, 10.6, 4.18, 10.6, 4.6, WHITE)
text(slide, "01", 12.25, 0.42, 0.45, 0.3, 11, MUTED, True, align=PP_ALIGN.RIGHT)

# 2. Why
slide = base("The operational question", "Context", 2)
text(slide, "How can a transport operator move from historical records to a live view of delay risk?", 0.8, 1.85, 7.5, 0.85, 25, WHITE, True, "Aptos Display")
bullets(slide, [
    "Historical CSV data provides realistic route, weather, congestion, and delay relationships.",
    "Synthetic events create a continuous stream without exposing real passenger data.",
    "The dashboard turns raw events into operational signals: delay rate, route volume, and risk.",
], 0.82, 3.05, 7.1, 2.4, 16, MUTED)
box(slide, 9.0, 2.0, 2.8, 2.8, NAVY_2, LINE)
text(slide, "FROM", 9.3, 2.35, 2.2, 0.2, 10, CYAN, True, align=PP_ALIGN.CENTER)
text(slide, "CSV rows", 9.3, 2.75, 2.2, 0.36, 23, WHITE, True, "Aptos Display", align=PP_ALIGN.CENTER)
arrow(slide, 10.4, 3.35, 10.4, 3.75, ORANGE)
text(slide, "TO", 9.3, 4.05, 2.2, 0.2, 10, ORANGE, True, align=PP_ALIGN.CENTER)
text(slide, "decisions", 9.3, 4.42, 2.2, 0.36, 23, WHITE, True, "Aptos Display", align=PP_ALIGN.CENTER)

# 3. Architecture
slide = base("One pipeline, five responsibilities", "Architecture", 3)
nodes = [("Reference CSV", "empirical seed", CYAN), ("Trip generator", "typed events", BLUE), ("Kafka", "route-keyed stream", ORANGE), ("MySQL", "operational store", MINT), ("Dashboard", "live decisions", CYAN)]
for i, (name, sub, color) in enumerate(nodes):
    x = 0.85 + i * 2.45
    box(slide, x, 2.45, 1.85, 1.25, NAVY_2, color)
    box(slide, x, 2.45, 0.08, 1.25, color, color, False)
    text(slide, name, x + 0.18, 2.73, 1.5, 0.25, 15, WHITE, True)
    text(slide, sub, x + 0.18, 3.14, 1.5, 0.2, 10, MUTED)
    if i < len(nodes) - 1: arrow(slide, x + 1.85, 3.08, x + 2.38, 3.08, MUTED)
text(slide, "Parallel analytics path", 0.9, 4.65, 2.3, 0.25, 11, ORANGE, True)
box(slide, 3.25, 4.45, 2.35, 0.72, NAVY_2, ORANGE)
text(slide, "HDFS historical store", 3.4, 4.68, 2.05, 0.2, 13, WHITE, True, align=PP_ALIGN.CENTER)
arrow(slide, 4.42, 4.42, 4.42, 3.8, ORANGE)
box(slide, 6.2, 4.45, 2.35, 0.72, NAVY_2, ORANGE)
text(slide, "Spark MLlib training", 6.35, 4.68, 2.05, 0.2, 13, WHITE, True, align=PP_ALIGN.CENTER)
arrow(slide, 5.6, 4.8, 6.15, 4.8, ORANGE)
box(slide, 9.15, 4.45, 2.35, 0.72, NAVY_2, ORANGE)
text(slide, "Delay classifier", 9.3, 4.68, 2.05, 0.2, 13, WHITE, True, align=PP_ALIGN.CENTER)
arrow(slide, 8.55, 4.8, 9.1, 4.8, ORANGE)

# 4. Generation
slide = base("Generated events keep the data believable", "Data generation", 4)
add_kpi(slide, "NDJSON", "portable event format", 0.85, 1.95, CYAN)
add_kpi(slide, "route_id", "Kafka partition key", 3.8, 1.95, ORANGE)
add_kpi(slide, "GMT+2", "Rwanda event time", 6.75, 1.95, MINT)
add_kpi(slide, "1–1,000", "API batch range", 9.7, 1.95, BLUE)
text(slide, "Each generated event samples a complete source row, then assigns a new trip ID and current schedule timestamps.", 0.9, 3.65, 7.5, 0.55, 20, WHITE, True, "Aptos Display")
bullets(slide, ["Preserves correlations between route, district, weather, congestion, and delay.", "Converts numeric fields before publishing so downstream systems can aggregate them.", "Supports both direct MySQL generation and Kafka-backed publishing."], 0.9, 4.45, 7.5, 1.6, 15, MUTED)
box(slide, 9.35, 3.6, 2.75, 2.1, NAVY_2, LINE)
text(slide, "EVENT", 9.65, 3.9, 2.15, 0.18, 10, CYAN, True, align=PP_ALIGN.CENTER)
text(slide, "trip_id\nroute_id\nweather\ncongestion\ndelayed", 9.65, 4.25, 2.15, 1.1, 17, WHITE, True, align=PP_ALIGN.CENTER)

# 5. Streaming
slide = base("Kafka carries live events with ordering in mind", "Streaming layer", 5)
text(slide, "The producer uses `route_id` as the message key.", 0.85, 1.85, 6.3, 0.4, 22, WHITE, True, "Aptos Display")
text(slide, "That gives consistent partition selection: trips on the same route remain ordered while different routes can process in parallel.", 0.88, 2.5, 6.0, 0.65, 16, MUTED)
for i in range(3):
    box(slide, 8.2, 2.0 + i * 1.05, 2.85, 0.72, NAVY_2, [CYAN, ORANGE, MINT][i])
    text(slide, f"partition {i}", 8.4, 2.23 + i * 1.05, 1.1, 0.2, 13, WHITE, True)
    text(slide, ["KG 7 Ave", "KN 5 Rd", "Nyabugogo"][i], 9.65, 2.23 + i * 1.05, 1.15, 0.2, 11, MUTED, align=PP_ALIGN.RIGHT)
box(slide, 1.0, 4.1, 5.9, 1.2, NAVY_2, ORANGE)
text(slide, "Reliability choices", 1.3, 4.35, 2.3, 0.2, 13, ORANGE, True)
text(slide, "acks=all  ·  idempotence  ·  delivery acknowledgement", 1.3, 4.73, 5.0, 0.22, 14, WHITE, True)
text(slide, "At-least-once processing is explicit and observable; the producer waits for broker acknowledgement before returning HTTP 201.", 7.1, 5.55, 4.9, 0.55, 14, MUTED)

# 6. MySQL
slide = base("MySQL is the operational source of truth", "Storage", 6)
text(slide, "The dashboard now reads and writes one runtime store: MySQL.", 0.85, 1.85, 8.2, 0.45, 23, WHITE, True, "Aptos Display")
add_kpi(slide, "InnoDB", "durable table engine", 0.9, 2.75, MINT)
add_kpi(slide, "utf8mb4", "full text compatibility", 3.9, 2.75, CYAN)
add_kpi(slide, "UPSERT", "safe repeatable loads", 6.9, 2.75, ORANGE)
add_kpi(slide, "MySQL", "dashboard backend", 9.9, 2.75, BLUE)
box(slide, 0.9, 4.5, 11.2, 1.25, NAVY_2, LINE)
text(slide, "transport_trip_events", 1.25, 4.8, 3.0, 0.3, 18, CYAN, True, "Aptos Display")
text(slide, "trip identity  ·  route  ·  coordinates  ·  schedule  ·  weather  ·  traffic  ·  delay outcome  ·  generated timestamp", 4.3, 4.76, 7.2, 0.38, 14, WHITE)
text(slide, "Historical SQLite rows can be replicated once; new runtime events go directly to MySQL.", 0.9, 6.25, 8.8, 0.3, 14, MUTED)

# 7. Dashboard
slide = base("The dashboard turns events into decisions", "Operational view", 7)
text(slide, "A two-second refresh loop exposes the current operational picture.", 0.85, 1.85, 7.0, 0.4, 22, WHITE, True, "Aptos Display")
for i, (label, value, color) in enumerate([("stored live events", "LIVE", CYAN), ("observed delay rate", "%", ORANGE), ("predicted delays", "RISK", MINT), ("top active route", "ROUTE", BLUE)]):
    x = 0.85 + i * 3.0
    box(slide, x, 2.7, 2.65, 1.15, NAVY_2, color)
    text(slide, value, x + 0.2, 2.93, 2.25, 0.35, 24, color, True, "Aptos Display")
    text(slide, label.upper(), x + 0.2, 3.48, 2.25, 0.17, 9, MUTED, True)
box(slide, 0.85, 4.35, 5.35, 1.35, NAVY_2, LINE)
text(slide, "LIVE TRANSIT MAP", 1.15, 4.62, 2.2, 0.2, 10, CYAN, True)
text(slide, "Origin coordinates · route markers · predicted risk", 1.15, 5.03, 4.5, 0.25, 14, WHITE, True)
box(slide, 6.65, 4.35, 5.35, 1.35, NAVY_2, LINE)
text(slide, "DECISION SUPPORT", 6.95, 4.62, 2.2, 0.2, 10, ORANGE, True)
text(slide, "Route volume · transport mix · priority alerts", 6.95, 5.03, 4.5, 0.25, 14, WHITE, True)

# 8. ML
slide = base("Spark MLlib adds a planning-time risk signal", "Analytics", 8)
text(slide, "The training path is designed to avoid target leakage.", 0.85, 1.85, 7.0, 0.4, 22, WHITE, True, "Aptos Display")
bullets(slide, ["Reference data is uploaded to HDFS as the historical source.", "Spark reads from HDFS and trains a delay classifier with planning/dispatch features.", "Actual arrival and departure delay fields are excluded from training features.", "Metrics and the pipeline model are saved back to HDFS for inspection and reuse."], 0.9, 2.65, 6.9, 2.5, 15, MUTED)
box(slide, 8.55, 2.25, 2.8, 0.75, NAVY_2, CYAN); text(slide, "HDFS", 8.55, 2.5, 2.8, 0.2, 16, CYAN, True, align=PP_ALIGN.CENTER)
arrow(slide, 9.95, 3.0, 9.95, 3.45, CYAN)
box(slide, 8.55, 3.55, 2.8, 0.75, NAVY_2, ORANGE); text(slide, "SPARK MLlib", 8.55, 3.8, 2.8, 0.2, 16, ORANGE, True, align=PP_ALIGN.CENTER)
arrow(slide, 9.95, 4.3, 9.95, 4.75, ORANGE)
box(slide, 8.55, 4.85, 2.8, 0.75, NAVY_2, MINT); text(slide, "risk model", 8.55, 5.1, 2.8, 0.2, 16, MINT, True, align=PP_ALIGN.CENTER)

# 9. Demo
slide = base("A reproducible local demo", "Runbook", 9)
text(slide, "The platform is designed to run on native Windows with explicit service boundaries.", 0.85, 1.85, 9.0, 0.45, 21, WHITE, True, "Aptos Display")
steps = [("01", "Configure", "Create .env with MySQL credentials"), ("02", "Start", "Run Django and Kafka services"), ("03", "Publish", "POST /api/trips/publish/"), ("04", "Observe", "Open the dashboard at localhost:8000")]
for i, (num, title, detail) in enumerate(steps):
    y = 2.75 + i * 0.82
    box(slide, 0.95, y, 0.7, 0.48, [CYAN, BLUE, ORANGE, MINT][i], [CYAN, BLUE, ORANGE, MINT][i])
    text(slide, num, 0.95, y + 0.14, 0.7, 0.15, 11, NAVY, True, align=PP_ALIGN.CENTER)
    text(slide, title, 1.95, y + 0.02, 1.35, 0.2, 15, WHITE, True)
    text(slide, detail, 3.45, y + 0.03, 5.6, 0.2, 14, MUTED)
box(slide, 9.4, 2.75, 2.4, 2.7, NAVY_2, LINE)
text(slide, "QUICK COMMAND", 9.68, 3.05, 1.85, 0.2, 9, CYAN, True, align=PP_ALIGN.CENTER)
text(slide, "python\nmanage.py\nrunserver", 9.68, 3.55, 1.85, 1.05, 20, WHITE, True, "Aptos Display", align=PP_ALIGN.CENTER)
text(slide, "Dashboard + APIs\nready for inspection", 9.68, 4.88, 1.85, 0.35, 11, MUTED, align=PP_ALIGN.CENTER)

# 10. Close
slide = base("What this project demonstrates", "Takeaways", 10)
text(slide, "A small but complete data engineering system can make the whole lifecycle visible.", 0.85, 1.9, 8.2, 0.55, 25, WHITE, True, "Aptos Display")
takeaways = [("01", "Generate", "Realistic synthetic events from empirical data"), ("02", "Stream", "Route-aware Kafka delivery with acknowledgement"), ("03", "Store", "MySQL as the single operational source of truth"), ("04", "Explain", "Dashboard metrics and Spark-based risk modeling")]
for i, (num, title, detail) in enumerate(takeaways):
    y = 3.0 + i * 0.72
    text(slide, num, 0.95, y, 0.45, 0.2, 12, [CYAN, BLUE, MINT, ORANGE][i], True)
    text(slide, title, 1.7, y, 1.4, 0.2, 15, WHITE, True)
    text(slide, detail, 3.25, y, 6.8, 0.22, 14, MUTED)
box(slide, 10.2, 2.8, 1.65, 1.65, BLUE, BLUE)
text(slide, "LIVE\nDATA\nLOOP", 10.2, 3.2, 1.65, 0.75, 18, WHITE, True, "Aptos Display", align=PP_ALIGN.CENTER)
text(slide, "Thank you", 10.05, 5.5, 1.95, 0.3, 20, CYAN, True, "Aptos Display", align=PP_ALIGN.CENTER)

prs.save(OUT)
print(OUT)